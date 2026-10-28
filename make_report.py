"""Fill the college PBL template (Report template.docx) for the hospital project.

Outputs report_final.docx (built from the template itself, so the front page and
styles match exactly) and report_final.md (plain text mirror).
"""

import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn

BASE = r"C:/Users/devis/Documents/Qoder/2026-10-28/a8ac5a0d/dsa-hospital-scheduling"
TEMPLATE = r"C:/Users/devis/Downloads/Report template.docx"
C_PATH = BASE + r"/hospital_system.c"

TITLE_FIELDS = {
    "Course Name": "Data Structures and Algorithms - II",
    "Course Code": "BCSE0401",
    "Faculty Name": "Mr Vikas Tripathi",
    "Assignment Type": "Individual Assignment",
    "Student(s) Name": "Devvrat Tripathi",
    "ERP ID": "250133150078",
    "Team/Group No.": "N/A (Individual Assignment)",
    "Problem Statement": "Designing an Algorithmically Efficient Data Management and "
                         "Processing System for an AI-Based Hospital Appointment Scheduling System",
    "SDG(s)": "SDG 3 - Good Health and Well-being",
    "Reporting Month": "Month 3 - Final",
    "Date of Submission": "",
}

SECTIONS = [
    {
        "heading": "2. Student Details / Team Details",
        "table": [
            ["S.No.", "Student Name", "ERP ID / Roll No.", "Role / Contribution"],
            ["1", "Devvrat Tripathi", "250133150078",
             "All five modules: algorithm selection, C coding, testing and report writing"],
        ],
    },
    {
        "heading": "3. Executive Summary",
        "paras": [
            "This assignment builds a hospital scheduling system with DSA-II concepts. My C "
            "program stores patients in a BST, queues them in a heap, finds shortest paths "
            "with Dijkstra, plans theatre time with knapsack DP, and packs free slots with "
            "backtracking. This gives fast lookup, urgent patients first, and better theatre use.",
        ],
    },
    {
        "heading": "4. Introduction",
        "paras": [
            "Hospitals handle hundreds of appointments every day. If records are stored in "
            "plain lists, every search takes time proportional to the number of patients, and "
            "theatre plans are often made by hand. This course teaches better tools: trees, "
            "graphs, dynamic programming, backtracking and advanced indexing. I used one tool "
            "from each unit in a single C program. The aim is to show that the right data "
            "structure turns a slow manual process into a fast and clear one.",
        ],
    },
    {
        "heading": "5. Problem Definition",
        "paras": [
            "Problem: manage hospital appointments and resources efficiently. Input: patients "
            "with ID, name, disease and severity (1 to 10), a campus map with 10 departments, "
            "8 pending surgeries with durations and priorities, and leftover appointment "
            "requests. Output: instant patient lookup, an urgency-ordered queue, the fastest "
            "route between departments, a maximum-priority theatre plan for an 8-hour day, and "
            "an exact packing of a free slot. Users are reception staff, doctors and "
            "administrators. Constraints: sample data, C language, terminal only, no external "
            "libraries. Success means O(log n) search, correct shortest paths and optimal "
            "theatre use. The work supports SDG 3 (Good Health and Well-being).",
        ],
    },
    {
        "heading": "6. Industry / Application Overview",
        "paras": [
            "Hospitals already use computers for this work. Big hospital systems store "
            "records in indexed databases, apps like Practo and Zocdoc book appointments, and "
            "ambulance services use shortest path searches for routing. These products are "
            "closed, so a student cannot see the data structures inside them. This project "
            "builds a small open version where each decision can be traced to one algorithm "
            "from the syllabus. That makes it useful for learning and easy to explain.",
        ],
    },
    {
        "heading": "7. Module 1: Tree-Based Data Management (BST and Heap)",
        "paras": [
            "This module uses binary search trees, heaps and heap sort (Unit 1). A BST "
            "keeps smaller IDs on the left and larger ones on the right, so each comparison "
            "removes half of the remaining records. A heap keeps every parent larger than its "
            "children, so the most urgent patient sits at the root. I stored records in a BST "
            "keyed by ID, the queue in a max-heap by severity, and used heap sort for ordered "
            "printing. Output: Arjun Rao (severity 9) prints first, Neha Gupta (2) last. "
            "Lesson: a plain BST turns into a linked list if IDs arrive in order, so an AVL "
            "tree is safer at scale.",
        ],
        "table": [
            ["Rank", "Patient", "Disease", "Severity"],
            ["1", "Arjun Rao", "Cardiac", "9"],
            ["2", "Priya Singh", "Fracture", "7"],
            ["3", "Sara Khan", "Asthma", "6"],
        ],
        "caption": "Figure 1: First three pops of the heap appointment queue.",
    },
    {
        "heading": "8. Module 2: Graph Modeling and Network Optimization (Dijkstra)",
        "paras": [
            "This module applies graph modelling and Dijkstra's shortest path algorithm "
            "(Unit 2) [1][4]. Dijkstra "
            "settles the nearest unvisited department first and relaxes its neighbours, which "
            "is correct because all corridor times are positive. I modelled the campus as a "
            "weighted undirected graph with 10 departments and 14 corridors in an adjacency "
            "matrix, then ran Dijkstra for the fastest route between any two departments. "
            "Output: ER to OT costs 7 minutes through ICU, better than 8 on the direct "
            "corridor. Lesson: an adjacency matrix is easy to read for a small campus; BFS "
            "and DFS check connectivity, and Prim and Kruskal both give a 34-minute spanning "
            "tree for cabling.",
        ],
        "table": [
            ["From", "To", "Cost", "Route"],
            ["ER", "OT", "7 min", "ER - ICU - OT"],
            ["Reception", "Cafeteria", "7 min", "Reception - Pharmacy - Cafeteria"],
        ],
        "caption": "Figure 2: Key shortest path results from the program.",
    },
    {
        "heading": "9. Module 3: Dynamic Programming for Optimization (0/1 Knapsack)",
        "paras": [
            "This module uses 0/1 knapsack dynamic programming (Unit 3). DP fits this problem "
            "because the same smaller questions repeat: the best plan for 480 minutes is "
            "built from the best plans for smaller capacities. I treated each surgery as an "
            "item with duration as weight and priority as value. The table dp[i][c] holds the "
            "best priority using the first i surgeries within c minutes, and tracing back "
            "gives the chosen list. Output: six surgeries with total priority 41, using all "
            "480 minutes. Greedy by priority per minute reaches only 39, so the DP plan wins. "
            "Lesson: DP runs in O(n x C) time and space. LCS fits symptom matching, resource "
            "allocation DP fits doctor distribution, and matrix chain multiplication does not "
            "fit this problem.",
        ],
        "table": [
            ["Surgery", "Duration (min)", "Priority"],
            ["Gallbladder removal", "120", "9"],
            ["Angioplasty", "150", "10"],
            ["Hernia repair", "75", "7"],
            ["Tonsillectomy", "60", "6"],
            ["Cataract", "45", "5"],
            ["Biopsy", "30", "4"],
        ],
        "caption": "Figure 3: Chosen surgeries from the DP plan (priority 41, 480 min used).",
    },
    {
        "heading": "10. Module 4: Backtracking and Branch and Bound (Sum of Subsets)",
        "paras": [
            "This module applies backtracking over a state space tree (Unit 4). The algorithm "
            "builds a "
            "solution one request at a time. At each step it tries to include the request, "
            "and it drops that branch as soon as the sum crosses the target, so all 32 "
            "subsets are not checked. I used it to fill a free slot of 80 minutes from "
            "leftover requests {30, 45, 60, 20, 35}, and it found {45, 35}. Lesson: the worst "
            "case is still exponential, but pruning keeps it fast for short lists. I also "
            "compared related problems: shift rostering like N-Queens, graph coloring to send "
            "overlapping surgeries to different theatres, and branch and bound with a reduced "
            "cost matrix for the ambulance route [2].",
        ],
        "figure": [
            "                 [ ] sum = 0",
            "                /         \\",
            "          take 45           skip 45",
            "          sum = 45          sum = 0",
            "         /        \\             ...",
            "    take 35      skip 35",
            "    sum = 80     sum = 45",
            "   FOUND {45, 35}",
        ],
        "caption": "Figure 4: Search tree for the 80-minute slot (pruned branches not shown).",
    },
    {
        "heading": "11. Module 5: Advanced Data Structures (B+ Tree)",
        "paras": [
            "This module studies B+ trees and advanced heaps (Unit 5) [5]. In a B+ tree all "
            "records sit "
            "in the leaves and internal nodes hold only separators; the leaves are linked, so "
            "a range query walks from one leaf to the next instead of returning to the root. "
            "I compared Red-Black trees, B-Trees, B+ trees, binomial heaps and Fibonacci "
            "heaps for a hospital storing lakhs of records. The B+ tree suits the patient "
            "index best because it matches disk blocks and handles date ranges. My worked "
            "example indexes 12 records by appointment day, and the query for days 8 to 20 "
            "returns 6 records in one leaf scan. For queues, Fibonacci heaps give O(1) "
            "amortized decrease-key for Dijkstra at large scale, while a binary heap is "
            "enough here.",
        ],
        "figure": [
            "                     [ 20 ]",
            "                   /        \\",
            "          [ 8 | 12 ]          [ 25 ]",
            "         /    |    \\          /    \\",
            "  [1 3 5]  [8 10]  [12 15 18]  [20 22]  [25 28]",
        ],
        "caption": "Figure 5: B+ tree indexing 12 records by appointment day (linked leaves).",
    },
    {
        "heading": "12. Final Integrated Solution / Testing and Validation",
        "paras": [
            "The C program ties all five modules together in one terminal menu. When a "
            "patient is added, the record goes into both the BST and the heap. The queue "
            "always pops the most urgent patient, Dijkstra answers route questions, the "
            "knapsack planner prints the theatre list, and subset search fills a free slot. "
            "I tested each option on its own first and then together, with sample data and "
            "some wrong inputs. Missing IDs return not found, bad department numbers are "
            "rejected, and duplicate IDs are refused. The five main test cases are listed in "
            "Figure 6 and all results matched hand calculations. The source is standard C99 "
            "with no external libraries, so it compiles with gcc or any online C compiler.",
        ],
        "table": [
            ["Test", "Input", "Expected", "Actual", "Status"],
            ["1. Search hit", "ID 105", "Arjun Rao found", "Found, severity 9", "Pass"],
            ["2. Search miss", "ID 999", "Not found", "Not found", "Pass"],
            ["3. Shortest path", "ER to OT", "7 min via ICU", "7 min via ICU", "Pass"],
            ["4. DP vs greedy", "480 min, 8 surgeries", "DP >= greedy", "41 vs 39", "Pass"],
            ["5. Slot packing", "Target 80", "Exact subset", "{45, 35}", "Pass"],
        ],
        "caption": "Figure 6: Test case summary.",
    },
    {
        "heading": "13. Expected Benefits / Learning Outcomes",
        "paras": [
            "Benefits: staff can find a patient in O(log n) steps instead of scanning a "
            "register, urgent cases move ahead automatically, and the theatre plan uses the "
            "full day at the highest possible priority. Learning: I now know when a tree is "
            "needed and when a heap is enough, why DP beats greedy here, how pruning makes "
            "backtracking practical, and why large systems index records with B+ trees. I "
            "also practised writing and explaining complexity analysis for each module.",
        ],
    },
    {
        "heading": "14. Challenges Faced",
        "paras": [
            "Some parts were harder than expected. Fitting five units into one small program "
            "forced me to pick one clear algorithm per unit instead of many half-done ones. I "
            "also learnt that a BST is fast only while it stays balanced, and my first heap "
            "test printed the wrong order because I compared severity in the wrong direction. "
            "I fixed it by checking the pop order against a sorted list. Fixed-size arrays "
            "and the missing GUI are known limits of this prototype, listed under future work.",
        ],
    },
    {
        "heading": "15. Future Improvements",
        "paras": [
            "Next steps: replace the BST with an AVL or B+ tree, save records to a file or "
            "database, predict severity with machine learning, use real corridor distances, "
            "and add a simple GUI.",
        ],
    },
    {
        "heading": "16. Conclusion",
        "paras": [
            "The project delivers a working hospital scheduling program in C that touches all "
            "five units of DSA-II. The BST registry and heap queue cover trees, Dijkstra "
            "handles routing, knapsack DP builds the theatre plan, backtracking packs free "
            "slots, and the report explains why a B+ tree is the right index at full scale. "
            "Every test passed. The biggest takeaway is that standard data structures solve "
            "real hospital problems well when each one is matched to the job it does best.",
        ],
    },
]

REFERENCES = [
    "T. H. Cormen, C. E. Leiserson, R. L. Rivest and C. Stein, Introduction to Algorithms, "
    "3rd ed. Cambridge, MA: MIT Press, 2009.",
    "E. Horowitz, S. Sahni and S. Rajasekaran, Fundamentals of Computer Algorithms, 2nd ed. "
    "Hyderabad: Universities Press, 2008.",
    "T. Cayirli and E. Veral, \"Outpatient scheduling in health care: A review of "
    "literature,\" Production and Operations Management, vol. 12, no. 4, pp. 519-549, 2003.",
    "GeeksforGeeks, \"Dijkstra's Shortest Path Algorithm,\" 2026. [Online]. Available: "
    "https://www.geeksforgeeks.org/dijkstras-shortest-path-algorithm-greedy-algo-7/. "
    "[Accessed: Oct. 28, 2026].",
    "GeeksforGeeks, \"Introduction of B+ Tree,\" 2026. [Online]. Available: "
    "https://www.geeksforgeeks.org/introduction-of-b-tree-2/. [Accessed: Oct. 28, 2026].",
    "NPTEL, \"Data Structures and Algorithms,\" IIT. [Online]. Available: "
    "https://nptel.ac.in/courses/106102064. [Accessed: Oct. 28, 2026].",
]

# ---------------- word count ----------------
def count_words(text):
    return len(re.findall(r"\S+", text))

limits = {"3": 50, "4": 75, "5": 100, "6": 75, "7": 100, "8": 100, "9": 150,
          "10": 150, "11": 150, "12": 200, "13": 100, "14": 100, "15": 50, "16": 100}
print("Section word counts (limit = template maximum):")
total = 0
for s in SECTIONS:
    num = s["heading"].split(".")[0]
    w = sum(count_words(p) for p in s.get("paras", []))
    total += w
    limit = limits.get(num)
    flag = ""
    if limit is not None:
        flag = "OK" if w <= limit * 1.1 else "OVER"
    print(f"  {s['heading'][:58]:60s} {w:4d} / {limit} {flag}")
print(f"  TOTAL body words (max 1500): {total}")

# ---------------- build docx from the template ----------------
doc = Document(TEMPLATE)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15

def find_para(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None

# fill front page table
title_table = doc.tables[0]
for row in title_table.rows[1:]:
    label = row.cells[0].text.strip()
    for key, value in TITLE_FIELDS.items():
        if label.startswith(key):
            cell = row.cells[1]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(12)
            break

# individual declaration
decl = find_para("I/We certify")
decl.text = ("I certify that this assignment is my own work, with all sources appropriately "
             "acknowledged. I understand that a 0% mark and further academic penalties may "
             "apply if the content is found to be plagiarised or inappropriately AI-generated.")

sig = find_para("Student Name & Signature(s):")
sig.text = ""
r = sig.add_run("Student Name & Signature(s): ")
r.bold = True
sig.add_run("Devvrat Tripathi")
sig.add_run().add_break(WD_BREAK.LINE)
r = sig.add_run("Date: ")
r.bold = True
sig.add_run("____________________")

# remove all guidance/instruction content after the declaration
start = None
for p in doc.paragraphs:
    if p.text.strip().startswith("PBL Report Structure"):
        start = p._p
        break
if start is not None:
    el = start
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt

# ---------------- add filled sections ----------------
def add_heading(text):
    p = doc.add_paragraph(style="Heading 3")
    p.add_run(text)

def add_para(text):
    doc.add_paragraph(text)

def add_table(rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(10.5)
            run.bold = (i == 0)

def add_mono(lines, size=9):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0

def add_caption(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)

for s in SECTIONS:
    add_heading(s["heading"])
    for para in s.get("paras", []):
        add_para(para)
    if s.get("table"):
        add_table(s["table"])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    if s.get("figure"):
        add_mono(s["figure"])
    if s.get("caption"):
        add_caption(s["caption"])

add_heading("17. References")
for i, ref in enumerate(REFERENCES, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"[{i}] ").bold = True
    p.add_run(ref)

doc.add_page_break()
add_heading("18. Appendix (Source Code)")
add_para("Source file: hospital_system.c. Compile: gcc hospital_system.c -o hospital_system, "
         "then run the executable. The program is standard C99 with no external libraries.")
with open(C_PATH, encoding="utf-8") as f:
    code_lines = f.read().rstrip("\n").split("\n")
add_mono(code_lines, size=8)

doc.save(BASE + r"/report_final.docx")

# ---------------- markdown mirror ----------------
md = ["# PBL Progress Report - Final", ""]
md.append("**Program and Branch Name:** B.Tech CSE (AI)")
md.append("")
md.append("| Field | Details |")
md.append("|---|---|")
md.append("| Course Name | Data Structures and Algorithms - II |")
md.append("| Course Code | BCSE0401 |")
md.append("| Faculty Name | Mr Vikas Tripathi |")
md.append("| Assignment Type | Individual Assignment |")
md.append("| Student(s) Name | Devvrat Tripathi |")
md.append("| ERP ID / Roll No. | 250133150078 |")
md.append("| Team/Group No. | N/A (Individual Assignment) |")
md.append("| Problem Statement / Project Title | Designing an Algorithmically Efficient Data "
          "Management and Processing System for an AI-Based Hospital Appointment Scheduling System |")
md.append("| SDG(s) | SDG 3 - Good Health and Well-being |")
md.append("| Reporting Month | Month 3 - Final |")
md.append("")
md.append("## Student Declaration")
md.append("")
md.append("I certify that this assignment is my own work, with all sources appropriately "
          "acknowledged. I understand that a 0% mark and further academic penalties may apply "
          "if the content is found to be plagiarised or inappropriately AI-generated.")
md.append("")
md.append("Student Name & Signature(s): Devvrat Tripathi  ")
md.append("Date: ____________________")
md.append("")
for s in SECTIONS:
    md.append(f"## {s['heading']}")
    md.append("")
    for para in s.get("paras", []):
        md.append(para)
        md.append("")
    if s.get("table"):
        rows = s["table"]
        md.append("| " + " | ".join(rows[0]) + " |")
        md.append("|" + "---|" * len(rows[0]))
        for row in rows[1:]:
            md.append("| " + " | ".join(row) + " |")
        md.append("")
    if s.get("figure"):
        md.append("```")
        md.extend(s["figure"])
        md.append("```")
        md.append("")
    if s.get("caption"):
        md.append("*" + s["caption"] + "*")
        md.append("")
md.append("## 17. References")
md.append("")
for i, ref in enumerate(REFERENCES, 1):
    md.append(f"{i}. {ref}")
md.append("")
md.append("## 18. Appendix (Source Code)")
md.append("")
md.append("Full file: hospital_system.c. Compile: `gcc hospital_system.c -o hospital_system`.")
md.append("")
md.append("```c")
md.extend(code_lines)
md.append("```")

with open(BASE + r"/report_final.md", "w", encoding="utf-8") as f:
    f.write("\n".join(md))

print("\nGenerated: report_final.docx (from template) and report_final.md")
